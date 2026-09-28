# utils/word_parser.py
import io
import zipfile
from xml.etree import ElementTree as ET

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None

try:
    import textract
except ImportError:  # pragma: no cover
    textract = None


def _extract_docx_text(file_bytes: bytes) -> str:
    """Extract readable text from a .docx byte stream."""
    if Document is None:
        raise RuntimeError("python-docx is not installed")

    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]

    table_rows = []
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                table_rows.append(" | ".join(cells))

    return "\n".join(paragraphs + table_rows).strip()


def _extract_docx_xml_fallback(file_bytes: bytes) -> str:
    """Fallback parser for DOCX files when the docx library is unavailable."""
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
            xml = zf.read("word/document.xml")
    except (zipfile.BadZipFile, KeyError):
        return ""

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    root = ET.fromstring(xml)
    paragraphs = []

    for para in root.findall(".//w:p", ns):
        texts = [node.text for node in para.findall(".//w:t", ns) if node.text]
        if texts:
            paragraphs.append("".join(texts))

    return "\n".join(paragraphs).strip()


def extract_word_text(file_bytes: bytes, filename: str = "document.docx") -> str:
    """Extract text from .docx / .doc input bytes, returning plain text."""
    name = (filename or "document.docx").lower()

    if name.endswith(".docx"):
        try:
            return _extract_docx_text(file_bytes)
        except Exception:
            fallback = _extract_docx_xml_fallback(file_bytes)
            if fallback:
                return fallback
            return ""

    if name.endswith(".doc"):
        if textract is not None:
            try:
                extracted = textract.process(io.BytesIO(file_bytes), extension="doc")
                if isinstance(extracted, bytes):
                    text = extracted.decode("utf-8", errors="ignore")
                    return text.strip()
                return str(extracted).strip()
            except Exception:
                pass

    return ""
