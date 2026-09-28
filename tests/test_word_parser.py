from io import BytesIO

from docx import Document

from utils.word_parser import extract_word_text


def test_extract_word_text_from_docx_bytes():
    doc = Document()
    doc.add_paragraph("Hello from Word support")
    doc.add_paragraph("Second paragraph")

    payload = BytesIO()
    doc.save(payload)
    payload.seek(0)

    text = extract_word_text(payload.getvalue(), filename="sample.docx")

    assert "Hello from Word support" in text
    assert "Second paragraph" in text
